import os
import pymupdf # PyMuPDF
import docx
import openpyxl
from PIL import Image
import pytesseract
import logging

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if ext == '.pdf':
            doc = pymupdf.open(file_path)
            for page in doc:
                text += page.get_text() + "\n"
        elif ext == '.docx':
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif ext == '.xlsx':
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_data = [str(cell) for cell in row if cell is not None]
                    if row_data:
                        text += " ".join(row_data) + "\n"
        elif ext == '.txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
        elif ext in ['.png', '.jpg', '.jpeg', '.webp']:
            # ប្រើ OCR សម្រាប់រូបភាព (គាំទ្រទាំងខ្មែរនិងអង់គ្លេស)
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img, lang='khm+eng')
        else:
            return None, f"❌ មិនគាំទ្រប្រភេទឯកសារ {ext} នេះទេ!"
            
        if not text.strip():
            return None, "❌ មិនមានអក្សរនៅក្នុងឯកសារនេះទេ!"
            
        return text.strip(), None
    except Exception as e:
        logging.error(f"Error extracting text from {file_path}: {e}")
        return None, f"❌ បរាជ័យក្នុងការអានឯកសារ៖ {str(e)}"

async def process_document(msg, bot):
    import os
    import asyncio
    
    file_obj = None
    if msg.photo:
        file_obj = await bot.get_file(msg.photo[-1].file_id)
        ext = ".jpg"
    elif msg.document:
        file_obj = await bot.get_file(msg.document.file_id)
        ext = os.path.splitext(msg.document.file_name)[1]
    else:
        return None, "❌ មិនមានឯកសារទេ"

    temp_path = f"temp_doc_{msg.message_id}{ext}"
    try:
        await file_obj.download_to_drive(temp_path)
        # ត្រូវប្រើ asyncio.to_thread ព្រោះ extract_text គឺមានដំណើរការយឺត (blocking)
        ex_txt, err = await asyncio.to_thread(extract_text, temp_path)
        return ex_txt, err
    except Exception as e:
        return None, f"❌ កំហុសពេលទាញយកឯកសារ៖ {str(e)}"
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except:
                pass
